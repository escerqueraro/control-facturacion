import os
import json
import sqlite3
import datetime
import uuid
from typing import List, Dict, Any, Optional
from fastapi import FastAPI, WebSocket, WebSocketDisconnect, HTTPException, Depends, UploadFile, File, Form
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import cloud_db

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
DB_PATH = os.path.join(BASE_DIR, "facturacion.db")

def get_db_connection():
    cloud_db.DB_PATH = DB_PATH
    return cloud_db.get_db()
os.makedirs(STATIC_DIR, exist_ok=True)
os.makedirs(UPLOADS_DIR, exist_ok=True)

app = FastAPI(title="Control de Facturación - Multi-Registro en Tiempo Real")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 41 Clients Master Seed Data
CLIENTS_MASTER = [
    {"id": 1, "name": "AIR FRANCE - BOGOTA - TRIPULACIONES", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "Portal API Global", "obs": "Se genera factura, se carga en API Global", "reqInf": False, "reqOc": False, "canal": "Portal API Global", "keyDay": "Primeros 5 días"},
    {"id": 2, "name": "ANDI", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "Marisol Marín", "obs": "Se envía informe a Marisol Marín, dada la respuesta se factura.", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 3, "name": "API", "prog": "* MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "API", "obs": "Se genera factura, primeros días del mes siguiente.", "reqInf": False, "reqOc": False, "canal": "Correo electrónico", "keyDay": "Días 1-5 mes siguiente"},
    {"id": 4, "name": "APTAR", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Cristina Rojas (cc Erika Zapata)", "obs": "Se envía informe a Cristina Rojas con copia a Erika Zapata, a vuelta de correo envían orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 02 y Día 16"},
    {"id": 5, "name": "AVIANCA - BOGOTA", "prog": "CADA 05-23", "freqType": "Quincenal", "resp": "Sol", "contact": "Edison Torres, A. Peralta, L. Ramirez", "obs": "Se envía informe a Edison Torres - Alejandro Peralta - Luis Ramirez y se espera orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 05 y Día 23"},
    {"id": 6, "name": "AVIANCA - CALI", "prog": "CADA 05-23", "freqType": "Quincenal", "resp": "Sol", "contact": "Edison Torres, A. Peralta, L. Ramirez", "obs": "Se envía informe a Edison Torres - Alejandro Peralta - Luis Ramirez y se espera orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 05 y Día 23"},
    {"id": 7, "name": "AVIANCA - EQUIPAJES", "prog": "SEMANAL", "freqType": "Semanal", "resp": "Sol", "contact": "Avianca Equipajes", "obs": "Se realiza informes semanales junto los vouchers de transporte y al finalizar el mes, se debe generar un resumen en el informe.", "reqInf": True, "reqOc": True, "canal": "Correo + Vouchers", "keyDay": "Semanal + Cierre mes"},
    {"id": 8, "name": "BANCO DE LA REPUBLICA", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Don Hector", "obs": "Se genera factura de inmediato, don Hector facilita la información.", "reqInf": False, "reqOc": False, "canal": "Factura electrónica", "keyDay": "Inmediato tras servicio"},
    {"id": 9, "name": "BANCO DE OCCIDENTE", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Autorizador de viaje", "obs": "Revisar el número de viaje y la información de quién autoriza.", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 10, "name": "BMS", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Mary Holguín, Paula Martinez y equipo", "obs": "Se realizan facturas, una por cada orden de compra. 1 solo informe con 1 hoja de Excel por OC.", "reqInf": True, "reqOc": True, "canal": "Correo con Excel x OC", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 11, "name": "CINEMARK COLOMBIA", "prog": "MES ANTICIPADO", "freqType": "Mensual", "resp": "Sol", "contact": "Alexander Victoria", "obs": "Se genera factura, mes anticipado y se envía factura a Alexander Victoria (más tardar a 5 del mes en curso)", "reqInf": False, "reqOc": False, "canal": "Correo electrónico", "keyDay": "Máximo día 05 mes curso"},
    {"id": 12, "name": "COLEGIO COLOMBO BRITANICO", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Administración CCB", "obs": "Se genera factura", "reqInf": False, "reqOc": False, "canal": "Factura electrónica", "keyDay": "Inmediato"},
    {"id": 13, "name": "COMFENALCO", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Fabian Campuzano", "obs": "Factura debe enviarse lunes a miércoles antes de los 20 del mes siguiente. Factura + Informe a Fabian Campuzano", "reqInf": True, "reqOc": False, "canal": "Correo (Factura + Inf)", "keyDay": "Lun-Mié antes día 20"},
    {"id": 14, "name": "COPA", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Zulay", "obs": "Se genera factura, luego de generar factura se envía factura + informe a Zulay.", "reqInf": True, "reqOc": False, "canal": "Correo (Factura + Inf)", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 15, "name": "EML", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Laura Vega", "obs": "Se envía la información a Laura Vega para aprobación de factura", "reqInf": True, "reqOc": True, "canal": "Correo aprobación", "keyDay": "Inmediato tras servicio"},
    {"id": 16, "name": "ESTELAR", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Sra. Luz Cano", "obs": "Se envía la información a la señora Luz Cano, debe ir por número de referencia", "reqInf": True, "reqOc": False, "canal": "Correo por Ref", "keyDay": "Inmediato tras servicio"},
    {"id": 17, "name": "EXPEDITORS", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Sol", "contact": "Patricia Delgado", "obs": "Luego de generar factura se envía factura + informe a Patricia Delgado", "reqInf": True, "reqOc": False, "canal": "Correo (Factura + Inf)", "keyDay": "Inmediato tras servicio"},
    {"id": 18, "name": "FEDERACION COLOMBIANA DE BALONCESTO", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Sra. Fanny Ceballos", "obs": "Se envía información a Fanny Ceballos, una vez apruebe pide link de pago", "reqInf": True, "reqOc": True, "canal": "Correo + Link pago", "keyDay": "Inmediato tras servicio"},
    {"id": 19, "name": "FLEISCHMANN", "prog": "CADA 16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Melissa Peña", "obs": "Se envía la información cada 16 a Melissa Peña, a vuelta de correo da orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 16 de cada mes"},
    {"id": 20, "name": "FONDO DE EMPLEADOS BANCO DE OCCIDENTE", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Kevin Mauricio Castañeda", "obs": "Se envía informe a Kevin Mauricio Castañeda", "reqInf": True, "reqOc": False, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 21, "name": "FUNDACION VALLE DEL LILI", "prog": "* INMEDIATO", "freqType": "Por Evento", "resp": "Edwar", "contact": "Solicitante del servicio", "obs": "Se envía el informe a quien solicito el servicio para que generen orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo a solicitante", "keyDay": "Inmediato tras servicio"},
    {"id": 22, "name": "GENFAR", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Cada solicitante", "obs": "Se envía informe a cada persona que solicita.", "reqInf": True, "reqOc": False, "canal": "Correo a solicitantes", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 23, "name": "GENFAR - PLANTA", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Katerine Arcila", "obs": "Se envía informe a Katerine Arcila, servicios de planta y adicionales", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 02 y Día 16"},
    {"id": 24, "name": "HENRY LUIS VILLEGAS", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Hector Mejía", "obs": "Se envía informe a Hector Mejía, luego se factura.", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 02 y Día 16"},
    {"id": 25, "name": "HOTEL CASA VALLECAUCANA", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Administración Hotel", "obs": "Se envía informe y luego de la respuesta se factura.", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 26, "name": "IN BOND GEMA", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "Fabian Soto", "obs": "Se envía la factura a Fabian Soto", "reqInf": False, "reqOc": False, "canal": "Correo a Fabian Soto", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 27, "name": "LATAM - BOGOTA", "prog": "RECEPCIÓN DE LATAM", "freqType": "Por Evento", "resp": "Edwar", "contact": "LATAM / Sistema GOW", "obs": "Se revisa información, modificaciones en GOW (si las hay) y se factura.", "reqInf": True, "reqOc": False, "canal": "Factura tras GOW", "keyDay": "Al recibir corte LATAM"},
    {"id": 28, "name": "LATAM - CALI", "prog": "CADA 15 - FIN DE MES", "freqType": "Quincenal", "resp": "Sol", "contact": "Cada Área", "obs": "Se envía informe cada área, se espera HES y se factura", "reqInf": True, "reqOc": True, "canal": "Factura tras HES", "keyDay": "Día 15 y Fin de mes"},
    {"id": 29, "name": "LATAM - EQUIPAJES", "prog": "CADA 15 - FIN DE MES", "freqType": "Quincenal", "resp": "Sol", "contact": "Cada Área", "obs": "Se envía informe cada área, se espera HES y se factura", "reqInf": True, "reqOc": True, "canal": "Factura tras HES", "keyDay": "Día 15 y Fin de mes"},
    {"id": 30, "name": "MENZIES", "prog": "CADA 23", "freqType": "Quincenal", "resp": "Sol", "contact": "Natalia Avendaño", "obs": "Se envía la factura a Natalia Avendaño", "reqInf": False, "reqOc": False, "canal": "Correo a Natalia", "keyDay": "Día 23 de cada mes"},
    {"id": 31, "name": "MINA SERVICIOS", "prog": "CADA 15", "freqType": "Quincenal", "resp": "Sol", "contact": "Natalia Jaramillo", "obs": "Se envía el informe a Natalia Jaramillo y genera orden de compra para facturar", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 15 de cada mes"},
    {"id": 32, "name": "PGI", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Maritzabel Barrezueta", "obs": "Se envía información a Maritzabel Barrezueta, si a 5 días no hay novedad se factura", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 02 y Día 16"},
    {"id": 33, "name": "PRODUCTOS RAMO", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "Contacto Ramo", "obs": "Se envía informe y una vez aprueben se factura", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 34, "name": "QBCO", "prog": "CADA 23", "freqType": "Quincenal", "resp": "Sol", "contact": "Contacto QBCO", "obs": "Se envía informe y una vez aprueben se factura", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Día 23 de cada mes"},
    {"id": 35, "name": "S.O.S.", "prog": "CADA 14", "freqType": "Quincenal", "resp": "Sol", "contact": "Betty Peréz", "obs": "Luego de generar factura se envía factura + informe a Betty Peréz", "reqInf": True, "reqOc": False, "canal": "Correo (Factura + Inf)", "keyDay": "Día 14 de cada mes"},
    {"id": 36, "name": "SANOFI", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Cada solicitante", "obs": "Se envía informe a cada persona que solicita.", "reqInf": True, "reqOc": False, "canal": "Correo a solicitantes", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 37, "name": "SPIRAX", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "Administración Spirax", "obs": "Se factura y se envía el PDF con el informe", "reqInf": True, "reqOc": False, "canal": "Correo (Factura PDF + Inf)", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 38, "name": "STF -  BOGOTA RUTA", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Edwar", "contact": "asistente.comercial@studiof.com.co", "obs": "Se envían informe al correo comercial y se espera orden de compra", "reqInf": True, "reqOc": True, "canal": "Correo electrónico", "keyDay": "Fin de mes / Días 1-5"},
    {"id": 39, "name": "STF -  CALI", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Víctor / Funcionaria STF", "obs": "STF envía 01 y 16 informe. Revisar con Víctor, aprobar y esperar entrada contable. Suministros es el 19% del total.", "reqInf": True, "reqOc": True, "canal": "Factura tras Entrada Contable", "keyDay": "Día 01-02 y Día 16"},
    {"id": 40, "name": "SUMINISTROS DE LA INDUSTRIA", "prog": "CADA 02-16", "freqType": "Quincenal", "resp": "Edwar", "contact": "Relacionado con STF Cali", "obs": "Facturación quincenal ligada a entrada contable de STF Cali (19% del valor total).", "reqInf": False, "reqOc": True, "canal": "Factura 19% STF", "keyDay": "Día 02 y Día 16"},
    {"id": 41, "name": "VIRUTEX", "prog": "MENSUAL", "freqType": "Mensual", "resp": "Sol", "contact": "Administración Virutex", "obs": "Se factura y se envía el PDF con el informe", "reqInf": True, "reqOc": False, "canal": "Correo (Factura PDF + Inf)", "keyDay": "Fin de mes / Días 1-5"}
]

# Database Setup & Migrations
def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    # 1. System Users Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS system_users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE,
        password TEXT,
        name TEXT,
        role TEXT DEFAULT 'admin',
        cargo TEXT DEFAULT 'Analista de Facturación',
        theme TEXT DEFAULT 'edwar',
        filter TEXT DEFAULT 'ALL',
        is_active INTEGER DEFAULT 1,
        created_at TEXT DEFAULT '',
        updated_at TEXT DEFAULT ''
    )
    """)

    # 2. Clients Table (Persistent Client Master)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS clients (
        id INTEGER PRIMARY KEY,
        name TEXT,
        prog TEXT,
        freq_type TEXT,
        resp TEXT,
        contact TEXT,
        obs TEXT,
        req_inf INTEGER DEFAULT 1,
        req_oc INTEGER DEFAULT 1,
        canal TEXT DEFAULT 'Correo electrónico',
        key_day TEXT DEFAULT '',
        is_blocked INTEGER DEFAULT 0
    )
    """)

    # 3. Cargos Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS cargos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT UNIQUE
    )
    """)

    # 4. Matrix Cuatrimestre Tracking Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS matrix_cells (
        client_id INTEGER,
        period_col TEXT,
        status TEXT DEFAULT 'PENDIENTE',
        invoice_num TEXT DEFAULT '',
        invoice_val TEXT DEFAULT '',
        notes TEXT DEFAULT '',
        updated_by TEXT DEFAULT '',
        updated_at TEXT DEFAULT '',
        PRIMARY KEY (client_id, period_col)
    )
    """)

    # 5. Primary Multi-Entry Billing Records Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS billing_records (
        record_id TEXT PRIMARY KEY,
        client_id INTEGER,
        client_name TEXT,
        freq_type TEXT DEFAULT 'Mensual',
        month TEXT DEFAULT 'Octubre',
        period_detail TEXT DEFAULT 'Mes Completo',
        period_key TEXT DEFAULT 'OCT_1Q',
        step1 BOOLEAN DEFAULT 0,
        step1_date TEXT DEFAULT '',
        pref_num TEXT DEFAULT '',
        pref_val TEXT DEFAULT '',
        step2 BOOLEAN DEFAULT 0,
        step2_oc TEXT DEFAULT '',
        step2_date TEXT DEFAULT '',
        oc_val TEXT DEFAULT '',
        step3 BOOLEAN DEFAULT 0,
        step3_fac TEXT DEFAULT '',
        step3_date TEXT DEFAULT '',
        fac_val TEXT DEFAULT '',
        step4 BOOLEAN DEFAULT 0,
        step4_date TEXT DEFAULT '',
        notes TEXT DEFAULT '',
        updated_by TEXT DEFAULT '',
        updated_at TEXT DEFAULT ''
    )
    """)
    conn.commit()

    # Dynamic Column Migrations (safe, non-destructive)
    def add_col_if_missing(table_name, col_name, col_type):
        try:
            if getattr(conn, 'is_pg', False):
                cursor.execute(f"ALTER TABLE {table_name} ADD COLUMN IF NOT EXISTS {col_name} {col_type}")
                conn.commit()
            else:
                cursor.execute(f"ALTER TABLE {table_name} ADD COLUMN {col_name} {col_type}")
                conn.commit()
        except Exception:
            if hasattr(conn, 'rollback'):
                conn.rollback()

    add_col_if_missing("billing_records", "area_group", "TEXT DEFAULT ''")
    add_col_if_missing("billing_records", "fac_pdf_url", "TEXT DEFAULT ''")
    add_col_if_missing("billing_records", "report_doc_url", "TEXT DEFAULT ''")
    add_col_if_missing("system_users", "permissions", "TEXT DEFAULT 'total'")

    # 6. File Archive Table (Persistent file storage across container restarts)
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS uploaded_files_archive (
        file_url TEXT PRIMARY KEY,
        filename TEXT,
        content_type TEXT,
        file_bytes BLOB,
        record_id TEXT,
        doc_type TEXT,
        created_at TEXT
    )
    """)
    conn.commit()

    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Seed Default Users if empty
    cursor.execute("SELECT COUNT(*) FROM system_users")
    if cursor.fetchone()[0] == 0:
        default_users = [
            ("master", ".soporte", "Administrador Master", "master", "Administrador Master del Sistema", "edwar", "ALL"),
            ("1143851865", "1143851865", "Sol", "admin", "Especialista de Facturación & Cartera", "sol", "Sol"),
            ("1107078586", "1107078586", "Edwar", "admin", "Coordinador de Facturación Operativa", "edwar", "Edwar"),
            ("1234", "1234", "Invitado", "viewer", "Consultor Externo / Modo Lectura", "guest", "ALL")
        ]
        for u, p, n, r, c, t, f in default_users:
            cursor.execute("""
            INSERT INTO system_users (username, password, name, role, cargo, theme, filter, is_active, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
            """, (u, p, n, r, c, t, f, now_iso, now_iso))
        conn.commit()

    # Ensure Master User exists even if table was pre-existing
    cursor.execute("SELECT COUNT(*) FROM system_users WHERE role = 'master'")
    if cursor.fetchone()[0] == 0:
        cursor.execute("""
        INSERT INTO system_users (username, password, name, role, cargo, theme, filter, is_active, created_at, updated_at)
        VALUES ('master', '.soporte', 'Administrador Master', 'master', 'Administrador Master del Sistema', 'edwar', 'ALL', 1, ?, ?)
        """, (now_iso, now_iso))
        conn.commit()

    # Seed Default Clients if empty
    cursor.execute("SELECT COUNT(*) FROM clients")
    if cursor.fetchone()[0] == 0:
        for c in CLIENTS_MASTER:
            cursor.execute("""
            INSERT INTO clients (id, name, prog, freq_type, resp, contact, obs, req_inf, req_oc, canal, key_day, is_blocked)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """, (c["id"], c["name"], c["prog"], c["freqType"], c["resp"], c["contact"], c["obs"],
                  1 if c["reqInf"] else 0, 1 if c["reqOc"] else 0, c["canal"], c["keyDay"]))
        conn.commit()

    # Seed Default Cargos if empty
    cursor.execute("SELECT COUNT(*) FROM cargos")
    if cursor.fetchone()[0] == 0:
        default_cargos = [
            "Administrador Master del Sistema",
            "Especialista de Facturación & Cartera",
            "Coordinador de Facturación Operativa",
            "Analista Senior de Facturación",
            "Analista Junior de Facturación",
            "Líder de Facturación & Cartera",
            "Asistente de Facturación",
            "Consultor Externo / Modo Lectura"
        ]
        for cargo in default_cargos:
            cursor.execute("INSERT OR IGNORE INTO cargos (title) VALUES (?)", (cargo,))
        conn.commit()

    # Seed Default Billing Records if empty
    cursor.execute("SELECT COUNT(*) FROM billing_records")
    if cursor.fetchone()[0] == 0:
        cursor.execute("SELECT * FROM clients ORDER BY id ASC")
        c_rows = cursor.fetchall()
        for c in c_rows:
            cid, cname, cprog, cfreq, cresp, ccontact, cobs, creq_inf, creq_oc, ccanal, ckey_day, cblocked = c
            rec_id = f"rec_{cid}_default"
            p_prog = cprog or ""
            if "02-16" in p_prog or "05-23" in p_prog or "15" in p_prog:
                f_type = "Quincenal"
                p_detail = "1Q"
            elif "SEMANAL" in p_prog:
                f_type = "Semanal"
                p_detail = "Semana 1"
            elif "INMEDIATO" in p_prog:
                f_type = "Por Evento"
                p_detail = "Servicio"
            else:
                f_type = "Mensual"
                p_detail = "Mes Completo"

            step1_def = 0
            step2_def = 0

            cursor.execute("""
            INSERT INTO billing_records
            (record_id, client_id, client_name, freq_type, month, period_detail, period_key, step1, step1_date, pref_num, pref_val, step2, step2_oc, step2_date, oc_val, step3, step3_fac, step3_date, fac_val, step4, step4_date, notes, updated_by, updated_at)
            VALUES (?, ?, ?, ?, 'Octubre', ?, 'OCT_1Q', ?, '', '', '', ?, '', '', '', 0, '', '', '', 0, '', '', 'Sistema', ?)
            """, (rec_id, cid, cname, f_type, p_detail, step1_def, step2_def, now_iso))
        conn.commit()

    # Seed Matrix Cells if empty
    cursor.execute("SELECT COUNT(*) FROM matrix_cells")
    if cursor.fetchone()[0] == 0:
        cursor.execute("SELECT id FROM clients")
        c_ids = [r[0] for r in cursor.fetchall()]
        cols = ["sep_1q", "sep_2q", "oct_1q", "oct_2q", "nov_1q", "nov_2q", "dic_1q", "dic_2q"]
        for cid in c_ids:
            for col in cols:
                # Default status
                cursor.execute("""
                INSERT OR IGNORE INTO matrix_cells (client_id, period_col, status, invoice_num, invoice_val, updated_by, updated_at)
                VALUES (?, ?, 'PENDIENTE', '', '', 'Sistema', ?)
                """, (cid, col, now_iso))
        conn.commit()

    # Automatic Cleanup of any residual test / mock data on startup
    cursor.execute("DELETE FROM clients WHERE id = 99 OR name LIKE '%TEST%' OR name LIKE '%RECUPERADA%'")
    cursor.execute("""
    DELETE FROM billing_records 
    WHERE client_id = 99 
       OR client_name LIKE '%TEST%' 
       OR client_name LIKE '%RECUPERADA%' 
       OR record_id LIKE '%test%'
       OR pref_num IN ('PF-RECOV-999', 'PF-AF-2026-01')
       OR step3_fac IN ('FAC-RECOV-789', 'FAC-AV-01', 'FAC-AV-02')
    """)
    cursor.execute("""
    DELETE FROM billing_records 
    WHERE client_id = 5 AND record_id != 'rec_5_default' AND (step3_fac LIKE 'FAC-AV%' OR fac_val LIKE '%15.000.000%' OR fac_val LIKE '%28.500.000%')
    """)
    cursor.execute("""
    UPDATE billing_records
    SET step1 = 0, step1_date = '', pref_num = '', pref_val = '',
        step2 = 0, step2_oc = '', step2_date = '', oc_val = '',
        step3 = 0, step3_fac = '', step3_date = '', fac_val = '',
        fac_pdf_url = '', report_doc_url = '',
        step4 = 0, step4_date = '', notes = ''
    WHERE record_id IN ('rec_1_default', 'rec_5_default') AND (pref_num = 'PF-AF-2026-01' OR step3_fac = 'FAC-AV-01' OR fac_val = '$ 4.500.000')
    """)
    cursor.execute("DELETE FROM uploaded_files_archive WHERE filename LIKE '%test%' OR record_id LIKE '%test%' OR record_id = 'rec_test_recovery'")
    conn.commit()

    conn.close()

init_db()

# WebSocket Connection Manager with Real-Time User Presence
class ConnectionManager:
    def __init__(self):
        self.active_connections: List[WebSocket] = []
        self.user_sockets: Dict[WebSocket, Dict[str, Any]] = {}

    async def connect(self, websocket: WebSocket):
        await websocket.accept()
        self.active_connections.append(websocket)

    def register_user(self, websocket: WebSocket, user_info: dict):
        self.user_sockets[websocket] = user_info

    def disconnect(self, websocket: WebSocket):
        if websocket in self.active_connections:
            self.active_connections.remove(websocket)
        if websocket in self.user_sockets:
            del self.user_sockets[websocket]

    def get_online_users(self) -> List[Dict[str, Any]]:
        seen = {}
        for u in self.user_sockets.values():
            un = u.get("username")
            if un:
                seen[un] = u
        return list(seen.values())

    async def broadcast(self, message: dict):
        dead_connections = []
        for connection in list(self.active_connections):
            try:
                await connection.send_json(message)
            except Exception:
                dead_connections.append(connection)
        for dc in dead_connections:
            self.disconnect(dc)

    async def broadcast_presence(self):
        try:
            conn = get_db_connection()
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            cursor.execute("SELECT id, username, name, role, cargo, is_active FROM system_users ORDER BY id ASC")
            users = [dict(r) for r in cursor.fetchall()]
            conn.close()
        except Exception:
            users = []

        online_usernames = {str(u.get("username", "")).strip().lower() for u in self.user_sockets.values() if u.get("username")}
        for u in users:
            u["is_online"] = str(u.get("username", "")).strip().lower() in online_usernames

        await self.broadcast({
            "type": "PRESENCE_UPDATE",
            "users": users,
            "online_users": self.get_online_users()
        })

manager = ConnectionManager()

# Data Models
class LoginRequest(BaseModel):
    password: str
    username: Optional[str] = None

class UpdateProfileRequest(BaseModel):
    username: str
    new_password: Optional[str] = None
    theme: Optional[str] = None

class CreateRecordRequest(BaseModel):
    client_id: int
    freq_type: str = "Quincenal"
    month: str = "Octubre"
    period_detail: str = "1Q"
    area_group: Optional[str] = ""
    updated_by: str = "Sistema"

class UpdateRecordFieldRequest(BaseModel):
    record_id: str
    field: str
    value: Any
    updated_by: str

class AdminUserCreateRequest(BaseModel):
    username: str
    password: str
    name: str
    role: str = "admin"
    cargo: str = "Analista de Facturación"
    theme: str = "edwar"
    filter: str = "ALL"
    permissions: Optional[str] = "total"

class AdminUserUpdateRequest(BaseModel):
    id: int
    username: str
    password: Optional[str] = None
    name: str
    role: str
    cargo: str
    theme: Optional[str] = "edwar"
    filter: Optional[str] = "ALL"
    permissions: Optional[str] = "total"
    is_active: Optional[int] = 1

class AdminClientCreateRequest(BaseModel):
    name: str
    prog: str
    freq_type: str
    resp: str
    contact: str
    obs: str
    req_inf: bool = True
    req_oc: bool = True
    canal: str = "Correo electrónico"
    key_day: str = ""

class AdminClientUpdateRequest(BaseModel):
    id: int
    name: str
    prog: str
    freq_type: str
    resp: str
    contact: str
    obs: str
    req_inf: bool
    req_oc: bool
    canal: str
    key_day: str
    is_blocked: int = 0

class AdminToggleBlockRequest(BaseModel):
    id: int

class CargoCreateRequest(BaseModel):
    title: str

class MatrixUpdateRequest(BaseModel):
    client_id: int
    period_col: str
    status: str
    invoice_num: str = ""
    invoice_val: str = ""
    updated_by: str = "Sistema"

# Endpoints
@app.post("/api/login")
def login(req: LoginRequest):
    pwd = (req.password or "").strip()
    usr = (req.username or "").strip()

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    user_row = None
    # 1. If username and password provided
    if usr and pwd:
        cursor.execute("SELECT * FROM system_users WHERE (username = ? OR password = ?) AND password = ? AND is_active = 1", (usr, pwd, pwd))
        user_row = cursor.fetchone()
    elif pwd:
        # Check by password
        cursor.execute("SELECT * FROM system_users WHERE password = ? AND is_active = 1", (pwd,))
        user_row = cursor.fetchone()
        if not user_row:
            # Check by username/cédula
            cursor.execute("SELECT * FROM system_users WHERE username = ? AND is_active = 1", (pwd,))
            user_row = cursor.fetchone()
    
    conn.close()

    if user_row:
        return {
            "success": True,
            "user": user_row["name"],
            "username": user_row["username"],
            "role": user_row["role"],
            "cargo": user_row["cargo"],
            "theme": user_row["theme"],
            "filter": user_row["filter"]
        }
    raise HTTPException(status_code=401, detail="Credenciales no autorizadas o usuario inactivo.")

@app.post("/api/user/update-profile")
async def update_profile(req: UpdateProfileRequest):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM system_users WHERE username = ?", (req.username,))
    user_row = cursor.fetchone()
    if not user_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Usuario no encontrado")

    new_pwd = user_row["password"]
    if req.new_password and req.new_password.strip():
        new_pwd = req.new_password.strip()

    new_theme = user_row["theme"]
    if req.theme and req.theme in ["sol", "edwar"]:
        if user_row["role"] != "viewer":
            new_theme = req.theme

    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute("UPDATE system_users SET password = ?, theme = ?, updated_at = ? WHERE username = ?",
                   (new_pwd, new_theme, now_iso, req.username))
    conn.commit()
    conn.close()

    return {
        "success": True,
        "username": req.username,
        "theme": new_theme,
        "message": "Perfil actualizado exitosamente"
    }

# Master Admin Users Management
@app.get("/api/admin/users")
def get_admin_users():
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, password, name, role, cargo, theme, filter, is_active, permissions FROM system_users ORDER BY id ASC")
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return {"users": rows}

@app.post("/api/admin/users/create")
def create_admin_user(req: AdminUserCreateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    perm = req.permissions if req.permissions else "total"
    try:
        cursor.execute("""
        INSERT INTO system_users (username, password, name, role, cargo, theme, filter, permissions, is_active, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)
        """, (req.username.strip(), req.password.strip(), req.name.strip(), req.role, req.cargo, req.theme, req.filter, perm, now_iso, now_iso))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        raise HTTPException(status_code=400, detail="El usuario o cédula ya existe en el sistema.")
    conn.close()
    return {"success": True, "message": "Usuario creado exitosamente"}

@app.post("/api/admin/users/update")
async def update_admin_user(req: AdminUserUpdateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    perm = req.permissions if req.permissions else "total"

    if req.password and req.password.strip():
        cursor.execute("""
        UPDATE system_users SET username = ?, password = ?, name = ?, role = ?, cargo = ?, theme = ?, filter = ?, permissions = ?, is_active = ?, updated_at = ?
        WHERE id = ?
        """, (req.username.strip(), req.password.strip(), req.name.strip(), req.role, req.cargo, req.theme, req.filter, perm, req.is_active, now_iso, req.id))
    else:
        cursor.execute("""
        UPDATE system_users SET username = ?, name = ?, role = ?, cargo = ?, theme = ?, filter = ?, permissions = ?, is_active = ?, updated_at = ?
        WHERE id = ?
        """, (req.username.strip(), req.name.strip(), req.role, req.cargo, req.theme, req.filter, perm, req.is_active, now_iso, req.id))

    conn.commit()
    conn.close()

    await manager.broadcast({"type": "USER_UPDATED", "user_id": req.id})
    await manager.broadcast_presence()
    return {"success": True, "message": "Usuario actualizado exitosamente"}

@app.delete("/api/admin/users/{user_id}")
def delete_admin_user(user_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT role FROM system_users WHERE id = ?", (user_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    if row[0] == "master":
        conn.close()
        raise HTTPException(status_code=400, detail="No se puede eliminar el usuario Master principal.")
    cursor.execute("DELETE FROM system_users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()
    return {"success": True, "message": "Usuario eliminado"}

# Master Admin Clients / Companies Management
@app.get("/api/admin/clients")
def get_admin_clients():
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM clients ORDER BY id ASC")
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return {"clients": rows}

@app.post("/api/admin/clients/create")
async def create_admin_client(req: AdminClientCreateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT MAX(id) FROM clients")
    max_id = cursor.fetchone()[0] or 0
    new_id = max_id + 1

    cursor.execute("""
    INSERT INTO clients (id, name, prog, freq_type, resp, contact, obs, req_inf, req_oc, canal, key_day, is_blocked)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
    """, (new_id, req.name.strip(), req.prog.strip(), req.freq_type, req.resp, req.contact.strip(), req.obs.strip(),
          1 if req.req_inf else 0, 1 if req.req_oc else 0, req.canal.strip(), req.key_day.strip()))
    conn.commit()

    # Also create default billing record
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rec_id = f"rec_{new_id}_default"
    p_detail = "1Q" if req.freq_type == "Quincenal" else ("Semana 1" if req.freq_type == "Semanal" else ("Servicio" if req.freq_type == "Por Evento" else "Mes Completo"))
    step1_def = 0 if req.req_inf else 1
    step2_def = 0 if req.req_oc else 1
    cursor.execute("""
    INSERT INTO billing_records
    (record_id, client_id, client_name, freq_type, month, period_detail, period_key, step1, step2, updated_by, updated_at)
    VALUES (?, ?, ?, ?, 'Octubre', ?, 'OCT_1Q', ?, ?, 'Master', ?)
    """, (rec_id, new_id, req.name.strip(), req.freq_type, p_detail, step1_def, step2_def, now_iso))

    # Also insert default matrix rows
    cols = ["sep_1q", "sep_2q", "oct_1q", "oct_2q", "nov_1q", "nov_2q", "dic_1q", "dic_2q"]
    for col in cols:
        cursor.execute("INSERT OR IGNORE INTO matrix_cells (client_id, period_col, status, updated_by, updated_at) VALUES (?, ?, 'PENDIENTE', 'Master', ?)", (new_id, col, now_iso))

    conn.commit()
    conn.close()

    await manager.broadcast({"type": "CLIENT_CREATED", "client_id": new_id, "name": req.name})
    return {"success": True, "id": new_id, "message": "Empresa creada exitosamente"}

@app.post("/api/admin/clients/update")
async def update_admin_client(req: AdminClientUpdateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
    UPDATE clients SET name = ?, prog = ?, freq_type = ?, resp = ?, contact = ?, obs = ?, req_inf = ?, req_oc = ?, canal = ?, key_day = ?, is_blocked = ?
    WHERE id = ?
    """, (req.name.strip(), req.prog.strip(), req.freq_type, req.resp, req.contact.strip(), req.obs.strip(),
          1 if req.req_inf else 0, 1 if req.req_oc else 0, req.canal.strip(), req.key_day.strip(), req.is_blocked, req.id))
    cursor.execute("UPDATE billing_records SET client_name = ? WHERE client_id = ?", (req.name.strip(), req.id))
    conn.commit()
    conn.close()

    await manager.broadcast({"type": "CLIENT_UPDATED", "client_id": req.id, "name": req.name})
    return {"success": True, "message": "Empresa actualizada"}

@app.post("/api/admin/clients/toggle-block")
async def toggle_block_client(req: AdminToggleBlockRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE clients SET is_blocked = CASE WHEN is_blocked = 1 THEN 0 ELSE 1 END WHERE id = ?", (req.id,))
    cursor.execute("SELECT is_blocked, name FROM clients WHERE id = ?", (req.id,))
    row = cursor.fetchone()
    conn.commit()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Empresa no encontrada")

    await manager.broadcast({"type": "CLIENT_BLOCKED_TOGGLED", "client_id": req.id, "is_blocked": row[0]})
    return {"success": True, "is_blocked": row[0], "name": row[1]}

@app.delete("/api/admin/clients/{client_id}")
async def delete_admin_client(client_id: int):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM clients WHERE id = ?", (client_id,))
    cursor.execute("DELETE FROM billing_records WHERE client_id = ?", (client_id,))
    cursor.execute("DELETE FROM matrix_cells WHERE client_id = ?", (client_id,))
    conn.commit()
    conn.close()

    await manager.broadcast({"type": "CLIENT_DELETED", "client_id": client_id})
    return {"success": True, "message": "Empresa eliminada"}

# Cargos Endpoints
@app.get("/api/cargos")
def get_cargos():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT title FROM cargos ORDER BY id ASC")
    cargos = [r[0] for r in cursor.fetchall()]
    conn.close()
    return {"cargos": cargos}

@app.post("/api/cargos/create")
def create_cargo(req: CargoCreateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO cargos (title) VALUES (?)", (req.title.strip(),))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        raise HTTPException(status_code=400, detail="Este cargo ya existe")
    conn.close()
    return {"success": True, "title": req.title.strip()}

# Matrix Endpoints
@app.get("/api/matrix")
def get_matrix_data():
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM clients ORDER BY id ASC")
    clients = [dict(c) for c in cursor.fetchall()]

    cursor.execute("SELECT * FROM matrix_cells")
    cells = cursor.fetchall()
    conn.close()

    matrix_map = {}
    for c in cells:
        cid = c["client_id"]
        if cid not in matrix_map:
            matrix_map[cid] = {}
        matrix_map[cid][c["period_col"]] = {
            "status": c["status"],
            "invoice_num": c["invoice_num"],
            "invoice_val": c["invoice_val"],
            "notes": c["notes"],
            "updated_by": c["updated_by"],
            "updated_at": c["updated_at"]
        }

    return {
        "clients": clients,
        "matrix": matrix_map
    }

@app.post("/api/matrix/update")
async def update_matrix_cell(req: MatrixUpdateRequest):
    conn = get_db_connection()
    cursor = conn.cursor()
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
    INSERT INTO matrix_cells (client_id, period_col, status, invoice_num, invoice_val, updated_by, updated_at)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    ON CONFLICT(client_id, period_col) DO UPDATE SET
    status = excluded.status,
    invoice_num = excluded.invoice_num,
    invoice_val = excluded.invoice_val,
    updated_by = excluded.updated_by,
    updated_at = excluded.updated_at
    """, (req.client_id, req.period_col, req.status, req.invoice_num, req.invoice_val, req.updated_by, now_iso))
    conn.commit()
    conn.close()

    broadcast_data = {
        "type": "MATRIX_CELL_UPDATED",
        "client_id": req.client_id,
        "period_col": req.period_col,
        "status": req.status,
        "invoice_num": req.invoice_num,
        "invoice_val": req.invoice_val,
        "updated_by": req.updated_by,
        "updated_at": now_iso
    }
    await manager.broadcast(broadcast_data)
    return {"success": True, "data": broadcast_data}

# Backward Compatibility Alias
@app.get("/api/clients")
def get_clients_alias(period: Optional[str] = None):
    res = get_records()
    res["clients"] = res["records"]
    return res

@app.post("/api/update")
async def update_alias(req: dict):
    record_id = req.get("record_id")
    client_id = req.get("client_id")
    field = req.get("field")
    value = req.get("value")
    updated_by = req.get("updated_by", "Sistema")
    
    if not record_id and client_id:
        conn = get_db_connection()
        c = conn.cursor()
        c.execute("SELECT record_id FROM billing_records WHERE client_id = ? ORDER BY record_id ASC LIMIT 1", (client_id,))
        row = c.fetchone()
        conn.close()
        if row:
            record_id = row[0]
            
    if not record_id:
        raise HTTPException(status_code=400, detail="record_id requerido")
        
    return await update_record_field(UpdateRecordFieldRequest(
        record_id=record_id,
        field=field,
        value=value,
        updated_by=updated_by
    ))

@app.get("/api/records")
def get_records(month: Optional[str] = None):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # Load clients from clients table (excluding any test companies)
    cursor.execute("SELECT * FROM clients WHERE id != 99 AND name NOT LIKE '%TEST%' ORDER BY id ASC")
    c_rows = cursor.fetchall()
    client_dict = {}
    for c in c_rows:
        client_dict[c["id"]] = {
            "id": c["id"],
            "name": c["name"],
            "prog": c["prog"],
            "freqType": c["freq_type"],
            "resp": c["resp"],
            "contact": c["contact"],
            "obs": c["obs"],
            "reqInf": bool(c["req_inf"]),
            "reqOc": bool(c["req_oc"]),
            "canal": c["canal"],
            "keyDay": c["key_day"],
            "is_blocked": bool(c["is_blocked"])
        }

    if month and month != "ALL":
        cursor.execute("SELECT * FROM billing_records WHERE month = ? AND client_id != 99 AND client_name NOT LIKE '%TEST%' AND record_id NOT LIKE '%test%' ORDER BY client_id ASC, record_id ASC", (month,))
    else:
        cursor.execute("SELECT * FROM billing_records WHERE client_id != 99 AND client_name NOT LIKE '%TEST%' AND record_id NOT LIKE '%test%' ORDER BY client_id ASC, record_id ASC")
    
    rows = cursor.fetchall()
    
    # Also fetch all rows to compute duplicate maps across the entire system
    cursor.execute("SELECT record_id, client_id, client_name, month, period_detail, pref_num, pref_val, step2_oc, oc_val, step3_fac, fac_val FROM billing_records WHERE client_id != 99 AND client_name NOT LIKE '%TEST%' AND record_id NOT LIKE '%test%'")
    all_rows = cursor.fetchall()
    conn.close()

    # Build Duplicate & Aggregation Maps
    # 1. Invoices Map
    invoices_map = {}
    for r in all_rows:
        fac = str(r["step3_fac"] or "").strip().upper()
        if fac:
            if fac not in invoices_map:
                invoices_map[fac] = []
            invoices_map[fac].append({
                "record_id": r["record_id"],
                "client_id": r["client_id"],
                "client_name": r["client_name"],
                "period": f"{r['month']} - {r['period_detail']}",
                "fac_val": r["fac_val"]
            })

    # 2. Pre-invoices Map
    preinvoices_map = {}
    for r in all_rows:
        pref = str(r["pref_num"] or "").strip().upper()
        if pref:
            if pref not in preinvoices_map:
                preinvoices_map[pref] = []
            preinvoices_map[pref].append({
                "record_id": r["record_id"],
                "client_id": r["client_id"],
                "client_name": r["client_name"],
                "period": f"{r['month']} - {r['period_detail']}",
                "pref_val": r["pref_val"]
            })

    # 3. OCs Map
    ocs_map = {}
    for r in all_rows:
        oc = str(r["step2_oc"] or "").strip().upper()
        if oc:
            if oc not in ocs_map:
                ocs_map[oc] = []
            ocs_map[oc].append({
                "record_id": r["record_id"],
                "client_id": r["client_id"],
                "client_name": r["client_name"],
                "period": f"{r['month']} - {r['period_detail']}",
                "fac_val": r["fac_val"],
                "oc_val": r["oc_val"]
            })

    # Merge client metadata
    records = []
    for r in rows:
        rec_dict = dict(r)
        c_meta = client_dict.get(rec_dict["client_id"], {})
        merged = {**c_meta, **rec_dict}
        records.append(merged)

    return {
        "records": records,
        "clients": list(client_dict.values()),
        "invoices_map": invoices_map,
        "preinvoices_map": preinvoices_map,
        "ocs_map": ocs_map
    }

@app.post("/api/records/create")
async def create_record(req: CreateRecordRequest):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM clients WHERE id = ?", (req.client_id,))
    c_row = cursor.fetchone()
    if not c_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Cliente no encontrado")

    c_meta = dict(c_row)
    new_id = f"rec_{req.client_id}_{uuid.uuid4().hex[:8]}"
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    step1_def = 0
    step2_def = 0

    area_group_val = req.area_group.strip() if req.area_group else ""
    cursor.execute("""
    INSERT INTO billing_records
    (record_id, client_id, client_name, freq_type, month, period_detail, period_key, area_group, step1, step1_date, pref_num, pref_val, step2, step2_oc, step2_date, oc_val, step3, step3_fac, step3_date, fac_val, step4, step4_date, notes, updated_by, updated_at)
    VALUES (?, ?, ?, ?, ?, ?, 'CUSTOM', ?, ?, '', '', '', ?, '', '', '', 0, '', '', '', 0, '', '', ?, ?)
    """, (new_id, req.client_id, c_meta["name"], req.freq_type, req.month, req.period_detail, area_group_val, step1_def, step2_def, req.updated_by, now_iso))
    conn.commit()
    conn.close()

    new_record = {
        "id": c_meta["id"],
        "name": c_meta["name"],
        "prog": c_meta["prog"],
        "freqType": c_meta["freq_type"],
        "resp": c_meta["resp"],
        "contact": c_meta["contact"],
        "obs": c_meta["obs"],
        "reqInf": bool(c_meta["req_inf"]),
        "reqOc": bool(c_meta["req_oc"]),
        "canal": c_meta["canal"],
        "keyDay": c_meta["key_day"],
        "is_blocked": bool(c_meta["is_blocked"]),
        "record_id": new_id,
        "client_id": req.client_id,
        "client_name": c_meta["name"],
        "freq_type": req.freq_type,
        "month": req.month,
        "period_detail": req.period_detail,
        "area_group": area_group_val,
        "fac_pdf_url": "",
        "report_doc_url": "",
        "step1": step1_def,
        "step1_date": "",
        "pref_num": "",
        "pref_val": "",
        "step2": step2_def,
        "step2_oc": "",
        "step2_date": "",
        "oc_val": "",
        "step3": 0,
        "step3_fac": "",
        "step3_date": "",
        "fac_val": "",
        "step4": 0,
        "step4_date": "",
        "notes": "",
        "updated_by": req.updated_by,
        "updated_at": now_iso
    }

    # Broadcast creation to all connected peers
    await manager.broadcast({
        "type": "RECORD_CREATED",
        "record": new_record,
        "author": req.updated_by
    })

    return {"success": True, "record": new_record}

@app.post("/api/records/update")
async def update_record_field(req: UpdateRecordFieldRequest):
    allowed_fields = [
        "freq_type", "month", "period_detail",
        "step1", "step1_date", "pref_num", "pref_val",
        "step2", "step2_oc", "step2_date", "oc_val",
        "step3", "step3_fac", "step3_date", "fac_val",
        "step4", "step4_date", "notes",
        "area_group", "fac_pdf_url", "report_doc_url"
    ]
    if req.field not in allowed_fields:
        raise HTTPException(status_code=400, detail="Campo no permitido")

    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = get_db_connection()
    cursor = conn.cursor()

    sql = f"UPDATE billing_records SET {req.field} = ?, updated_by = ?, updated_at = ? WHERE record_id = ?"
    cursor.execute(sql, (req.value, req.updated_by, now_iso, req.record_id))
    conn.commit()
    conn.close()

    # Broadcast field update to all connected peers
    broadcast_data = {
        "type": "RECORD_FIELD_UPDATED",
        "record_id": req.record_id,
        "field": req.field,
        "value": req.value,
        "updated_by": req.updated_by,
        "updated_at": now_iso
    }
    await manager.broadcast(broadcast_data)

    return {"success": True, "data": broadcast_data}

class CommitAndResetRequest(BaseModel):
    record_id: str
    client_id: int
    author: Optional[str] = "Sistema"
    freq_type: Optional[str] = None
    month: Optional[str] = None
    period_detail: Optional[str] = None
    area_group: Optional[str] = None
    step1: Optional[bool] = None
    step1_date: Optional[str] = None
    pref_num: Optional[str] = None
    pref_val: Optional[str] = None
    step2: Optional[bool] = None
    step2_oc: Optional[str] = None
    step2_date: Optional[str] = None
    oc_val: Optional[str] = None
    step3: Optional[bool] = None
    step3_fac: Optional[str] = None
    step3_date: Optional[str] = None
    fac_val: Optional[str] = None
    step4: Optional[bool] = None
    step4_date: Optional[str] = None
    notes: Optional[str] = None
    fac_pdf_url: Optional[str] = None
    report_doc_url: Optional[str] = None

@app.post("/api/records/commit_and_reset")
async def commit_and_reset_record(req: CommitAndResetRequest):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM billing_records WHERE record_id = ?", (req.record_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Registro no encontrado")

    current_data = dict(row)
    for field_name in [
        "freq_type", "month", "period_detail", "area_group",
        "step1", "step1_date", "pref_num", "pref_val",
        "step2", "step2_oc", "step2_date", "oc_val",
        "step3", "step3_fac", "step3_date", "fac_val",
        "step4", "step4_date", "notes",
        "fac_pdf_url", "report_doc_url"
    ]:
        val = getattr(req, field_name, None)
        if val is not None:
            current_data[field_name] = val

    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Generate unique ID for this committed record
    new_record_id = f"rec_{req.client_id}_{uuid.uuid4().hex[:8]}"

    # Insert permanent consolidated record
    cursor.execute("""
    INSERT INTO billing_records
    (record_id, client_id, client_name, freq_type, month, period_detail, period_key, area_group,
     step1, step1_date, pref_num, pref_val, step2, step2_oc, step2_date, oc_val,
     step3, step3_fac, step3_date, fac_val, fac_pdf_url, report_doc_url, step4, step4_date,
     notes, updated_by, updated_at)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        new_record_id, current_data["client_id"], current_data["client_name"],
        current_data["freq_type"], current_data["month"], current_data["period_detail"], current_data["period_key"],
        current_data.get("area_group", ""),
        current_data["step1"], current_data["step1_date"], current_data["pref_num"], current_data["pref_val"],
        current_data["step2"], current_data["step2_oc"], current_data["step2_date"], current_data["oc_val"],
        current_data["step3"], current_data["step3_fac"], current_data["step3_date"], current_data["fac_val"],
        current_data.get("fac_pdf_url", ""), current_data.get("report_doc_url", ""),
        current_data["step4"], current_data["step4_date"],
        current_data["notes"], req.author, now_iso
    ))

    # Reset the active input record (req.record_id) to empty/clean state!
    cursor.execute("""
    UPDATE billing_records
    SET step1 = 0, step1_date = '', pref_num = '', pref_val = '',
        step2 = 0, step2_oc = '', step2_date = '', oc_val = '',
        step3 = 0, step3_fac = '', step3_date = '', fac_val = '',
        fac_pdf_url = '', report_doc_url = '',
        step4 = 0, step4_date = '', notes = '',
        updated_by = ?, updated_at = ?
    WHERE record_id = ?
    """, (req.author, now_iso, req.record_id))

    conn.commit()

    cursor.execute("SELECT * FROM billing_records WHERE record_id = ?", (new_record_id,))
    committed_row = dict(cursor.fetchone())

    cursor.execute("SELECT * FROM billing_records WHERE record_id = ?", (req.record_id,))
    clean_row = dict(cursor.fetchone())

    conn.close()

    broadcast_data = {
        "type": "RECORD_COMMITTED_AND_RESET",
        "client_id": req.client_id,
        "committed_record": committed_row,
        "clean_record": clean_row,
        "author": req.author
    }
    await manager.broadcast(broadcast_data)

    return {
        "success": True,
        "committed_record": committed_row,
        "clean_record": clean_row
    }

class UpdateFullRecordRequest(BaseModel):
    record_id: str
    freq_type: Optional[str] = "Mensual"
    month: Optional[str] = "Octubre"
    period_detail: Optional[str] = "Mes Completo"
    area_group: Optional[str] = ""
    step1: Optional[bool] = False
    step1_date: Optional[str] = ""
    pref_num: Optional[str] = ""
    pref_val: Optional[str] = ""
    step2: Optional[bool] = False
    step2_oc: Optional[str] = ""
    step2_date: Optional[str] = ""
    oc_val: Optional[str] = ""
    step3: Optional[bool] = False
    step3_fac: Optional[str] = ""
    step3_date: Optional[str] = ""
    fac_val: Optional[str] = ""
    step4: Optional[bool] = False
    step4_date: Optional[str] = ""
    notes: Optional[str] = ""
    fac_pdf_url: Optional[str] = ""
    report_doc_url: Optional[str] = ""
    updated_by: Optional[str] = "Sistema"

@app.post("/api/records/update_full")
async def update_record_full(req: UpdateFullRecordRequest):
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("""
    UPDATE billing_records
    SET freq_type = ?, month = ?, period_detail = ?, area_group = ?,
        step1 = ?, step1_date = ?, pref_num = ?, pref_val = ?,
        step2 = ?, step2_oc = ?, step2_date = ?, oc_val = ?,
        step3 = ?, step3_fac = ?, step3_date = ?, fac_val = ?,
        step4 = ?, step4_date = ?, notes = ?,
        fac_pdf_url = ?, report_doc_url = ?,
        updated_by = ?, updated_at = ?
    WHERE record_id = ?
    """, (
        req.freq_type, req.month, req.period_detail, req.area_group,
        1 if req.step1 else 0, req.step1_date, req.pref_num, req.pref_val,
        1 if req.step2 else 0, req.step2_oc, req.step2_date, req.oc_val,
        1 if req.step3 else 0, req.step3_fac, req.step3_date, req.fac_val,
        1 if req.step4 else 0, req.step4_date, req.notes,
        req.fac_pdf_url, req.report_doc_url,
        req.updated_by, now_iso, req.record_id
    ))
    conn.commit()

    cursor.execute("SELECT * FROM billing_records WHERE record_id = ?", (req.record_id,))
    row = cursor.fetchone()
    conn.close()

    if not row:
        raise HTTPException(status_code=404, detail="Registro no encontrado")

    updated_record = dict(row)
    await manager.broadcast({
        "type": "RECORD_UPDATED_FULL",
        "record": updated_record,
        "author": req.updated_by
    })

    return {"success": True, "record": updated_record}

@app.delete("/api/records/{record_id}")
async def delete_record(record_id: str, user: str = "Sistema"):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT client_id, client_name FROM billing_records WHERE record_id = ?", (record_id,))
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise HTTPException(status_code=404, detail="Registro no encontrado")

    client_id = row[0]
    cursor.execute("SELECT COUNT(*) FROM billing_records WHERE client_id = ?", (client_id,))
    cnt = cursor.fetchone()[0]

    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if cnt <= 1:
        # If it's the only row for this client, reset it to empty/clean rather than throwing an error
        cursor.execute("""
        UPDATE billing_records
        SET step1 = 0, step1_date = '', pref_num = '', pref_val = '',
            step2 = 0, step2_oc = '', step2_date = '', oc_val = '',
            step3 = 0, step3_fac = '', step3_date = '', fac_val = '',
            fac_pdf_url = '', report_doc_url = '',
            step4 = 0, step4_date = '', notes = '', area_group = '',
            updated_by = ?, updated_at = ?
        WHERE record_id = ?
        """, (user, now_iso, record_id))
        conn.commit()

        cursor.execute("SELECT * FROM billing_records WHERE record_id = ?", (record_id,))
        conn.row_factory = sqlite3.Row
        reset_row = dict(cursor.fetchone())
        conn.close()

        await manager.broadcast({
            "type": "RECORD_RESET_CLEAN",
            "record": reset_row,
            "record_id": record_id,
            "client_id": client_id,
            "author": user
        })
        return {"success": True, "action": "reset_clean", "record": reset_row}

    cursor.execute("DELETE FROM billing_records WHERE record_id = ?", (record_id,))
    conn.commit()
    conn.close()

    await manager.broadcast({
        "type": "RECORD_DELETED",
        "record_id": record_id,
        "client_id": client_id,
        "author": user
    })

    return {"success": True, "action": "deleted", "deleted_id": record_id}

@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try:
        while True:
            raw_text = await websocket.receive_text()
            if raw_text == "ping":
                await websocket.send_text("pong")
                continue
            try:
                msg = json.loads(raw_text)
                if isinstance(msg, dict):
                    mtype = msg.get("type")
                    if mtype == "AUTH_CONNECT":
                        user_info = msg.get("user", {})
                        manager.register_user(websocket, user_info)
                        await manager.broadcast_presence()
                    elif mtype == "AUTH_LOGOUT":
                        manager.disconnect(websocket)
                        await manager.broadcast_presence()
            except Exception:
                pass
    except WebSocketDisconnect:
        manager.disconnect(websocket)
        await manager.broadcast_presence()
    except Exception:
        manager.disconnect(websocket)
        await manager.broadcast_presence()

@app.get("/api/presence/status")
def get_presence_status():
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, name, role, cargo, permissions, is_active FROM system_users ORDER BY id ASC")
    users = [dict(r) for r in cursor.fetchall()]
    conn.close()

    online_usernames = {str(u.get("username", "")).strip().lower() for u in manager.user_sockets.values() if u.get("username")}
    for u in users:
        u["is_online"] = str(u.get("username", "")).strip().lower() in online_usernames

    return {"users": users, "online_count": len(online_usernames)}

# File Upload and Document Management
import re

ALLOWED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".docx", ".doc", ".csv", ".png", ".jpg", ".jpeg"}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # 25 MB

@app.post("/api/files/upload")
async def upload_file(
    file: UploadFile = File(...),
    record_id: str = Form(...),
    doc_type: str = Form(...)  # 'fac_pdf' or 'report_doc'
):
    if doc_type not in ["fac_pdf", "report_doc"]:
        raise HTTPException(status_code=400, detail="Tipo de documento inválido ('fac_pdf' o 'report_doc').")

    filename = file.filename or "archivo"
    _, ext = os.path.splitext(filename)
    ext = ext.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Extensión no permitida ({ext}). Permitidas: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    clean_base = re.sub(r'[^a-zA-Z0-9_\-\.]', '_', os.path.splitext(filename)[0])[:35]
    unique_name = f"{doc_type}_{record_id}_{uuid.uuid4().hex[:6]}_{clean_base}{ext}"
    file_path = os.path.join(UPLOADS_DIR, unique_name)

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=400, detail="El archivo excede el tamaño máximo permitido de 25MB.")

    with open(file_path, "wb") as f:
        f.write(content)

    cloud_url = await cloud_db.upload_file_to_cloud(content, unique_name, file.content_type or 'application/octet-stream')
    file_url = cloud_url if cloud_url else f"/uploads/{unique_name}"
    col_name = "fac_pdf_url" if doc_type == "fac_pdf" else "report_doc_url"
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = get_db_connection()
    cursor = conn.cursor()
    # Save binary into uploaded_files_archive table for persistent disaster recovery
    cursor.execute("""
    INSERT OR REPLACE INTO uploaded_files_archive
    (file_url, filename, content_type, file_bytes, record_id, doc_type, created_at)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (file_url, unique_name, file.content_type or 'application/octet-stream', content, record_id, doc_type, now_iso))

    if doc_type == "report_doc":
        cursor.execute(f"UPDATE billing_records SET {col_name} = ?, step1 = 1, updated_at = ? WHERE record_id = ?", (file_url, now_iso, record_id))
    elif doc_type == "fac_pdf":
        cursor.execute(f"UPDATE billing_records SET {col_name} = ?, step3 = 1, updated_at = ? WHERE record_id = ?", (file_url, now_iso, record_id))
    else:
        cursor.execute(f"UPDATE billing_records SET {col_name} = ?, updated_at = ? WHERE record_id = ?", (file_url, now_iso, record_id))
    conn.commit()
    conn.close()

    broadcast_data = {
        "type": "RECORD_FIELD_UPDATED",
        "record_id": record_id,
        "field": col_name,
        "value": file_url,
        "updated_by": "Sistema (Archivo)",
        "updated_at": now_iso
    }
    await manager.broadcast(broadcast_data)

    if doc_type == "report_doc":
        await manager.broadcast({
            "type": "RECORD_FIELD_UPDATED",
            "record_id": record_id,
            "field": "step1",
            "value": True,
            "updated_by": "Sistema (Archivo)",
            "updated_at": now_iso
        })
    elif doc_type == "fac_pdf":
        await manager.broadcast({
            "type": "RECORD_FIELD_UPDATED",
            "record_id": record_id,
            "field": "step3",
            "value": True,
            "updated_by": "Sistema (Archivo)",
            "updated_at": now_iso
        })

    return {
        "success": True,
        "file_url": file_url,
        "filename": filename,
        "record_id": record_id,
        "doc_type": doc_type
    }

class DeleteFileRequest(BaseModel):
    record_id: str
    doc_type: str  # 'fac_pdf' or 'report_doc'

@app.post("/api/files/delete")
async def delete_file(req: DeleteFileRequest):
    if req.doc_type not in ["fac_pdf", "report_doc"]:
        raise HTTPException(status_code=400, detail="Tipo de documento inválido")

    col_name = "fac_pdf_url" if req.doc_type == "fac_pdf" else "report_doc_url"
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(f"SELECT {col_name} FROM billing_records WHERE record_id = ?", (req.record_id,))
    row = cursor.fetchone()
    old_url = row[0] if row else ""

    cursor.execute(f"UPDATE billing_records SET {col_name} = '', updated_at = ? WHERE record_id = ?", (now_iso, req.record_id))
    if old_url:
        cursor.execute("DELETE FROM uploaded_files_archive WHERE file_url = ? OR filename = ?", (old_url, os.path.basename(old_url)))
    conn.commit()
    conn.close()

    if old_url and old_url.startswith("/uploads/"):
        old_filename = os.path.basename(old_url)
        old_path = os.path.join(UPLOADS_DIR, old_filename)
        if os.path.exists(old_path):
            try:
                os.remove(old_path)
            except Exception:
                pass

    broadcast_data = {
        "type": "RECORD_FIELD_UPDATED",
        "record_id": req.record_id,
        "field": col_name,
        "value": "",
        "updated_by": "Sistema (Archivo eliminado)",
        "updated_at": now_iso
    }
    await manager.broadcast(broadcast_data)
    return {"success": True, "record_id": req.record_id, "doc_type": req.doc_type}

# Backup, Restore & LocalStorage Sync Endpoints
@app.get("/api/backup/export")
def export_backup():
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM clients")
    clients = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM billing_records")
    records = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM matrix_cells")
    matrix = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM cargos")
    cargos = [r["title"] for r in cursor.fetchall()]

    cursor.execute("SELECT id, username, password, name, role, cargo, theme, filter, is_active FROM system_users")
    users = [dict(r) for r in cursor.fetchall()]

    conn.close()

    backup_data = {
        "version": "1.0",
        "exported_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "clients": clients,
        "billing_records": records,
        "matrix_cells": matrix,
        "cargos": cargos,
        "users": users
    }
    return backup_data

@app.post("/api/backup/import")
async def import_backup(data: dict):
    if not isinstance(data, dict):
        raise HTTPException(status_code=400, detail="Formato de respaldo inválido")

    conn = get_db_connection()
    cursor = conn.cursor()
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Restore clients
    if "clients" in data and isinstance(data["clients"], list):
        for c in data["clients"]:
            cursor.execute("""
            INSERT OR REPLACE INTO clients (id, name, prog, freq_type, resp, contact, obs, req_inf, req_oc, canal, key_day, is_blocked)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (c.get("id"), c.get("name"), c.get("prog"), c.get("freq_type"), c.get("resp"), c.get("contact"), c.get("obs"),
                  1 if c.get("req_inf") else 0, 1 if c.get("req_oc") else 0, c.get("canal", ""), c.get("key_day", ""), 1 if c.get("is_blocked") else 0))

    # Restore billing records
    if "billing_records" in data and isinstance(data["billing_records"], list):
        for r in data["billing_records"]:
            cursor.execute("""
            INSERT OR REPLACE INTO billing_records
            (record_id, client_id, client_name, freq_type, month, period_detail, period_key, area_group, step1, step1_date, pref_num, pref_val, step2, step2_oc, step2_date, oc_val, step3, step3_fac, step3_date, fac_val, fac_pdf_url, report_doc_url, step4, step4_date, notes, updated_by, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                r.get("record_id"), r.get("client_id"), r.get("client_name"), r.get("freq_type"), r.get("month"), r.get("period_detail"),
                r.get("period_key", "OCT_1Q"), r.get("area_group", ""),
                1 if r.get("step1") else 0, r.get("step1_date", ""), r.get("pref_num", ""), r.get("pref_val", ""),
                1 if r.get("step2") else 0, r.get("step2_oc", ""), r.get("step2_date", ""), r.get("oc_val", ""),
                1 if r.get("step3") else 0, r.get("step3_fac", ""), r.get("step3_date", ""), r.get("fac_val", ""),
                r.get("fac_pdf_url", ""), r.get("report_doc_url", ""),
                1 if r.get("step4") else 0, r.get("step4_date", ""), r.get("notes", ""), r.get("updated_by", "Restauración"), now_iso
            ))

    # Restore matrix cells
    if "matrix_cells" in data and isinstance(data["matrix_cells"], list):
        for m in data["matrix_cells"]:
            cursor.execute("""
            INSERT OR REPLACE INTO matrix_cells (client_id, period_col, status, invoice_num, invoice_val, notes, updated_by, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                m.get("client_id"), m.get("period_col"), m.get("status", "PENDIENTE"),
                m.get("invoice_num", ""), m.get("invoice_val", ""), m.get("notes", ""),
                m.get("updated_by", "Restauración"), now_iso
            ))

    # Restore cargos
    if "cargos" in data and isinstance(data["cargos"], list):
        for title in data["cargos"]:
            cursor.execute("INSERT OR IGNORE INTO cargos (title) VALUES (?)", (title,))

    # Restore users (preserve master)
    if "users" in data and isinstance(data["users"], list):
        for u in data["users"]:
            if u.get("role") == "master":
                cursor.execute("UPDATE system_users SET password = ? WHERE role = 'master'", (u.get("password"),))
            else:
                cursor.execute("""
                INSERT OR REPLACE INTO system_users (id, username, password, name, role, cargo, theme, filter, is_active, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    u.get("id"), u.get("username"), u.get("password"), u.get("name"), u.get("role", "admin"),
                    u.get("cargo", "Analista"), u.get("theme", "edwar"), u.get("filter", "ALL"),
                    1 if u.get("is_active", 1) else 0, now_iso, now_iso
                ))

    conn.commit()
    conn.close()

    await manager.broadcast({"type": "DATABASE_RESTORED"})
    return {"success": True, "message": "Base de datos restaurada correctamente"}

@app.post("/api/sync/client_cache")
async def sync_client_cache(payload: dict):
    records = payload.get("records", [])
    if not records:
        return {"success": True, "synced": 0}

    conn = get_db_connection()
    cursor = conn.cursor()
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    synced_count = 0
    for r in records:
        rec_id = r.get("record_id")
        if not rec_id:
            continue
        cid = r.get("client_id")
        cname = str(r.get("client_name", "")).upper()
        pref_num = str(r.get("pref_num", ""))
        step3_fac = str(r.get("step3_fac", ""))

        # STRICT BLOCK: NEVER insert test or recovery dummy records!
        if cid == 99 or "TEST" in cname or "RECUPERADA" in cname or "test" in str(rec_id).lower() or pref_num == "PF-RECOV-999" or step3_fac in ["FAC-RECOV-789", "FAC-AV-01", "FAC-AV-02"]:
            continue

        cursor.execute("""
        INSERT OR REPLACE INTO billing_records
        (record_id, client_id, client_name, freq_type, month, period_detail, period_key, area_group, step1, step1_date, pref_num, pref_val, step2, step2_oc, step2_date, oc_val, step3, step3_fac, step3_date, fac_val, fac_pdf_url, report_doc_url, step4, step4_date, notes, updated_by, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            rec_id, r.get("client_id"), r.get("client_name"), r.get("freq_type", "Mensual"), r.get("month", "Octubre"),
            r.get("period_detail", "Mes Completo"), r.get("period_key", "OCT_1Q"), r.get("area_group", ""),
            1 if r.get("step1") else 0, r.get("step1_date", ""), r.get("pref_num", ""), r.get("pref_val", ""),
            1 if r.get("step2") else 0, r.get("step2_oc", ""), r.get("step2_date", ""), r.get("oc_val", ""),
            1 if r.get("step3") else 0, r.get("step3_fac", ""), r.get("step3_date", ""), r.get("fac_val", ""),
            r.get("fac_pdf_url", ""), r.get("report_doc_url", ""),
            1 if r.get("step4") else 0, r.get("step4_date", ""), r.get("notes", ""), r.get("updated_by", "AutoSync"), now_iso
        ))
        synced_count += 1

    conn.commit()
    conn.close()
    return {"success": True, "synced": synced_count}

@app.post("/api/admin/clean_test_data")
async def clean_test_data():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM clients WHERE id = 99 OR name LIKE '%TEST%' OR name LIKE '%RECUPERADA%'")
    cursor.execute("""
    DELETE FROM billing_records 
    WHERE client_id = 99 
       OR client_name LIKE '%TEST%' 
       OR client_name LIKE '%RECUPERADA%' 
       OR record_id LIKE '%test%'
       OR pref_num IN ('PF-RECOV-999', 'PF-AF-2026-01')
       OR step3_fac IN ('FAC-RECOV-789', 'FAC-AV-01', 'FAC-AV-02')
    """)
    cursor.execute("""
    DELETE FROM billing_records 
    WHERE client_id = 5 AND record_id != 'rec_5_default' AND (step3_fac LIKE 'FAC-AV%' OR fac_val LIKE '%15.000.000%' OR fac_val LIKE '%28.500.000%')
    """)
    cursor.execute("""
    UPDATE billing_records
    SET step1 = 0, step1_date = '', pref_num = '', pref_val = '',
        step2 = 0, step2_oc = '', step2_date = '', oc_val = '',
        step3 = 0, step3_fac = '', step3_date = '', fac_val = '',
        fac_pdf_url = '', report_doc_url = '',
        step4 = 0, step4_date = '', notes = ''
    WHERE record_id IN ('rec_1_default', 'rec_5_default')
    """)
    cursor.execute("DELETE FROM uploaded_files_archive WHERE filename LIKE '%test%' OR record_id LIKE '%test%' OR record_id = 'rec_test_recovery'")
    conn.commit()
    conn.close()
    await manager.broadcast({"type": "DATABASE_RESTORED"})
    return {"success": True, "message": "Datos de prueba purgados completamente"}

# Dynamic Excel Export
@app.get("/api/export/excel")
def export_excel(month: Optional[str] = "Octubre"):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    if month and month != "ALL":
        cursor.execute("SELECT * FROM billing_records WHERE month = ? ORDER BY client_id ASC, record_id ASC", (month,))
    else:
        cursor.execute("SELECT * FROM billing_records ORDER BY client_id ASC, record_id ASC")
    rows = cursor.fetchall()

    cursor.execute("SELECT * FROM clients")
    c_rows = cursor.fetchall()
    conn.close()

    client_dict = {c["id"]: dict(c) for c in c_rows}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Control_{month or 'General'}"

    NAVY = "1E3A8A"
    STEP1_COL = "0284C7"
    STEP2_COL = "D97706"
    STEP3_COL = "059669"
    STEP4_COL = "7C3AED"

    ws.merge_cells("A1:W1")
    ws["A1"] = f"CONTROL OPERATIVO DE FACTURACIÓN E INFORMES - {month or 'GENERAL'}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        ("ID", NAVY), ("CLIENTE", NAVY), ("RESPONSABLE", NAVY), ("PROGRAMACIÓN BASE", NAVY),
        ("TIPO CORTE", NAVY), ("MES", NAVY), ("PERÍODO / DETALLE", NAVY),
        ("P1: REQ. INFORME", STEP1_COL), ("P1: FECHA ENVÍO", STEP1_COL), ("P1: ¿ENVIADO?", STEP1_COL),
        ("P1: Nº PRE-FACTURA", STEP1_COL), ("P1: VALOR PRE-FACTURA", STEP1_COL),
        ("P2: REQ. OC", STEP2_COL), ("P2: Nº OC / HES", STEP2_COL), ("P2: FECHA OC", STEP2_COL), ("P2: ¿OC RECIBIDA?", STEP2_COL),
        ("P3: FECHA FACTURA", STEP3_COL), ("P3: Nº FACTURA", STEP3_COL), ("P3: VALOR FACTURA", STEP3_COL), ("P3: ¿GENERADA?", STEP3_COL),
        ("COMPARATIVA / AUDITORÍA", "D97706"),
        ("P4: CANAL ENTREGA", STEP4_COL), ("P4: ¿ENTREGADA?", STEP4_COL),
        ("ESTADO ACTUAL", "475569")
    ]

    for col_idx, (text, color) in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=text)
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    for r_idx, r in enumerate(rows, start=4):
        c_meta = client_dict.get(r["client_id"], {})
        s1 = "SÍ" if r["step1"] else ("N/A" if not c_meta.get("req_inf") else "NO")
        s2 = "SÍ" if r["step2"] else ("N/A" if not c_meta.get("req_oc") else "NO")
        s3 = "SÍ" if r["step3"] else "NO"
        s4 = "SÍ" if r["step4"] else "NO"

        # Math Difference
        p_val_str = str(r["pref_val"] or "").replace("$", "").replace(".", "").replace(",", "").strip()
        f_val_str = str(r["fac_val"] or "").replace("$", "").replace(".", "").replace(",", "").strip()
        try:
            p_val = float(p_val_str) if p_val_str else 0.0
            f_val = float(f_val_str) if f_val_str else 0.0
            diff = f_val - p_val
            diff_text = "✓ Coincide" if (p_val > 0 and f_val > 0 and diff == 0) else (f"Dif: ${diff:,.0f}" if (p_val > 0 or f_val > 0) else "")
        except Exception:
            diff_text = ""

        # Status
        if r["step4"]:
            st_text = "🟢 COMPLETADO"
        elif r["step3"]:
            st_text = "🔵 FACTURADO (Pend. Envío)"
        elif (r["step2"] or not c_meta.get("req_oc")) and (r["step1"] or not c_meta.get("req_inf")):
            st_text = "🟠 APROBADO (Pend. Facturar)"
        elif r["step1"]:
            st_text = "🟡 ESPERANDO OC/APROB."
        else:
            st_text = "⚪ PENDIENTE INFORME"

        row_vals = [
            r["client_id"], r["client_name"], c_meta.get("resp", ""), c_meta.get("prog", ""),
            r["freq_type"], r["month"], r["period_detail"],
            "SÍ" if c_meta.get("req_inf") else "NO", r["step1_date"], s1,
            r["pref_num"], r["pref_val"],
            "SÍ" if c_meta.get("req_oc") else "NO", r["step2_oc"], r["step2_date"], s2,
            r["step3_date"], r["step3_fac"], r["fac_val"], s3,
            diff_text,
            c_meta.get("canal", ""), s4,
            st_text
        ]

        fill_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.border = border_thin
            cell.font = Font(name="Calibri", size=10)
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            if col_idx in [1, 3, 4, 5, 6, 7, 8, 10, 13, 16, 20, 23]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Control_Facturacion_{month}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Cloud Health & Persistence Status
@app.get("/api/cloud/status")
def get_cloud_status_endpoint():
    return cloud_db.get_cloud_status()

# Monthly Closing Report Analysis
@app.get("/api/reports/monthly_closing")
def get_monthly_closing_report(month: Optional[str] = "Octubre"):
    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    if month and month != "ALL":
        cursor.execute("SELECT * FROM billing_records WHERE month = ? ORDER BY client_id ASC", (month,))
    else:
        cursor.execute("SELECT * FROM billing_records ORDER BY client_id ASC")
    records = cursor.fetchall()

    cursor.execute("SELECT * FROM clients")
    clients = {c["id"]: dict(c) for c in cursor.fetchall()}
    conn.close()

    total_records = len(records)
    pref_total = 0.0
    fac_total = 0.0

    step1_done = 0
    step2_done = 0
    step3_done = 0
    step4_done = 0

    bottlenecks = []
    by_resp = {}

    today = datetime.date.today()

    def parse_money(val_str):
        if not val_str:
            return 0.0
        cleaned = re.sub(r'[^\d]', '', str(val_str))
        try:
            return float(cleaned) if cleaned else 0.0
        except Exception:
            return 0.0

    def get_days_diff(date_str):
        if not date_str:
            return None
        try:
            d = datetime.datetime.strptime(date_str[:10], "%Y-%m-%d").date()
            return (today - d).days
        except Exception:
            return None

    for r in records:
        cid = r["client_id"]
        c_meta = clients.get(cid, {})
        resp = c_meta.get("resp", "Sin asignar")
        req_oc = bool(c_meta.get("req_oc", 1))
        req_inf = bool(c_meta.get("req_inf", 1))

        p_val = parse_money(r["pref_val"])
        f_val = parse_money(r["fac_val"])
        pref_total += p_val
        fac_total += f_val

        if resp not in by_resp:
            by_resp[resp] = {
                "resp": resp,
                "total": 0,
                "completed": 0,
                "invoiced": 0,
                "waiting_oc": 0,
                "pending": 0,
                "pref_sum": 0.0,
                "fac_sum": 0.0
            }
        by_resp[resp]["total"] += 1
        by_resp[resp]["pref_sum"] += p_val
        by_resp[resp]["fac_sum"] += f_val

        # Step counts
        if r["step1"]: step1_done += 1
        if r["step2"] or not req_oc: step2_done += 1
        if r["step3"]: step3_done += 1
        if r["step4"]:
            step4_done += 1
            by_resp[resp]["completed"] += 1
        elif r["step3"]:
            by_resp[resp]["invoiced"] += 1
        elif r["step1"] and req_oc and not r["step2"]:
            by_resp[resp]["waiting_oc"] += 1
        else:
            by_resp[resp]["pending"] += 1

        # Check bottlenecks
        if r["step3"] and not r["step4"]:
            days = get_days_diff(r["step3_date"]) or 0
            if days > 2:
                bottlenecks.append({
                    "record_id": r["record_id"],
                    "client_id": cid,
                    "client_name": r["client_name"],
                    "resp": resp,
                    "stage": "Radicación Pendiente",
                    "days_waiting": days,
                    "severity": "critical" if days > 5 else "warning",
                    "amount": f_val or p_val,
                    "details": f"Factura {r['step3_fac'] or 'emitida'} lleva {days} días sin entrega/radicación."
                })
        elif r["step1"] and req_oc and not r["step2"]:
            days = get_days_diff(r["step1_date"]) or 0
            if days > 3:
                bottlenecks.append({
                    "record_id": r["record_id"],
                    "client_id": cid,
                    "client_name": r["client_name"],
                    "resp": resp,
                    "stage": "Esperando Orden de Compra",
                    "days_waiting": days,
                    "severity": "critical" if days > 7 else "warning",
                    "amount": p_val,
                    "details": f"Informe enviado hace {days} días. Cliente aún no expide OC/HES."
                })
        elif (r["step2"] or not req_oc) and (r["step1"] or not req_inf) and not r["step3"]:
            days = get_days_diff(r["step2_date"] or r["step1_date"]) or 0
            if days > 2:
                bottlenecks.append({
                    "record_id": r["record_id"],
                    "client_id": cid,
                    "client_name": r["client_name"],
                    "resp": resp,
                    "stage": "Facturación Demorada",
                    "days_waiting": days,
                    "severity": "critical" if days > 4 else "warning",
                    "amount": p_val,
                    "details": f"Aprobación recibida hace {days} días. Pendiente emitir factura contable."
                })

    diff_val = pref_total - fac_total
    closing_rate = round((step4_done / total_records * 100), 1) if total_records > 0 else 0.0

    return {
        "success": True,
        "month": month,
        "total_records": total_records,
        "completed_records": step4_done,
        "invoiced_records": step3_done,
        "closing_rate_pct": closing_rate,
        "totals": {
            "pref_val": pref_total,
            "fac_val": fac_total,
            "diff_val": diff_val,
            "pref_val_fmt": f"${pref_total:,.0f}".replace(",", "."),
            "fac_val_fmt": f"${fac_total:,.0f}".replace(",", "."),
            "diff_val_fmt": f"${diff_val:,.0f}".replace(",", ".")
        },
        "funnel": {
            "step1_pct": round((step1_done / total_records * 100), 1) if total_records else 0,
            "step2_pct": round((step2_done / total_records * 100), 1) if total_records else 0,
            "step3_pct": round((step3_done / total_records * 100), 1) if total_records else 0,
            "step4_pct": round((step4_done / total_records * 100), 1) if total_records else 0,
            "step1_count": step1_done,
            "step2_count": step2_done,
            "step3_count": step3_done,
            "step4_count": step4_done
        },
        "by_resp": list(by_resp.values()),
        "bottlenecks": sorted(bottlenecks, key=lambda x: (0 if x["severity"] == "critical" else 1, -x["days_waiting"])),
        "bottlenecks_count": len(bottlenecks)
    }

# Export Monthly Closing Executive Excel
@app.get("/api/reports/export_closing_excel")
def export_closing_excel(month: Optional[str] = "Octubre"):
    closing_data = get_monthly_closing_report(month)
    wb = openpyxl.Workbook()

    NAVY = "1E3A8A"
    STEP1_COL = "0284C7"
    STEP2_COL = "D97706"
    STEP3_COL = "059669"
    STEP4_COL = "7C3AED"
    GRAY_BG = "F1F5F9"
    BORDER_COLOR = "CBD5E1"

    border_thin = Border(left=Side(style='thin', color=BORDER_COLOR), right=Side(style='thin', color=BORDER_COLOR),
                         top=Side(style='thin', color=BORDER_COLOR), bottom=Side(style='thin', color=BORDER_COLOR))

    # --- SHEET 1: RESUMEN EJECUTIVO DE CIERRE ---
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo Cierre"

    ws1.merge_cells("A1:G1")
    ws1["A1"] = f"TABLERO EJECUTIVO DE CIERRE MENSUAL OPERATIVO - {month.upper()}"
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws1["A1"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    # KPI summary cards
    kpis = [
        ("TOTAL REGISTROS", f"{closing_data['total_records']}", "475569"),
        ("TOTAL PREFACTURADO", f"{closing_data['totals']['pref_val_fmt']}", STEP1_COL),
        ("TOTAL FACTURADO REAL", f"{closing_data['totals']['fac_val_fmt']}", STEP3_COL),
        ("BRECHA / DIFERENCIA", f"{closing_data['totals']['diff_val_fmt']}", "DC2626" if closing_data['totals']['diff_val'] > 0 else "059669"),
        ("% AVANCE DE CIERRE", f"{closing_data['closing_rate_pct']}%", STEP4_COL)
    ]

    ws1.cell(row=3, column=1, value="INDICADORES CLAVE DE DESEMPEÑO (KPIs)").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for idx, (title, val, col) in enumerate(kpis, start=1):
        cell_t = ws1.cell(row=4, column=idx, value=title)
        cell_t.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell_t.fill = PatternFill(start_color=col, end_color=col, fill_type="solid")
        cell_t.alignment = Alignment(horizontal="center", vertical="center")

        cell_v = ws1.cell(row=5, column=idx, value=val)
        cell_v.font = Font(name="Calibri", size=13, bold=True, color="0F172A")
        cell_v.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        cell_v.alignment = Alignment(horizontal="center", vertical="center")
        cell_v.border = border_thin

    # Responsibility breakdown
    ws1.cell(row=7, column=1, value="CONCILIACIÓN POR RESPONSABLE ASIGNADO").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    resp_headers = ["Responsable", "Clientes Asignados", "Total Prefacturado", "Total Facturado Real", "Completados", "Pendientes", "% Éxito"]
    for c_idx, h_text in enumerate(resp_headers, start=1):
        c = ws1.cell(row=8, column=c_idx, value=h_text)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    curr_row = 9
    for resp_item in closing_data["by_resp"]:
        rate = round((resp_item["completed"] / resp_item["total"] * 100), 1) if resp_item["total"] else 0
        vals = [
            resp_item["resp"],
            resp_item["total"],
            f"${resp_item['pref_sum']:,.0f}".replace(",", "."),
            f"${resp_item['fac_sum']:,.0f}".replace(",", "."),
            resp_item["completed"],
            resp_item["total"] - resp_item["completed"],
            f"{rate}%"
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws1.cell(row=curr_row, column=c_idx, value=v)
            cell.font = Font(name="Calibri", size=10, bold=(c_idx == 1 or c_idx == 7))
            cell.alignment = Alignment(horizontal="center" if c_idx > 1 else "left", vertical="center")
            cell.border = border_thin
        curr_row += 1

    # Bottlenecks / Semáforo Crítico
    curr_row += 1
    ws1.cell(row=curr_row, column=1, value="SEMÁFORO DE CUELLOS DE BOTELLA Y ALERTAS CRÍTICAS").font = Font(name="Calibri", size=11, bold=True, color="B91C1C")
    curr_row += 1

    bot_headers = ["Nivel Alerta", "Cliente", "Responsable", "Fase Atascada", "Días Espera", "Valor en Riesgo", "Diagnóstico Operativo"]
    for c_idx, h_text in enumerate(bot_headers, start=1):
        c = ws1.cell(row=curr_row, column=c_idx, value=h_text)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
    curr_row += 1

    if closing_data["bottlenecks"]:
        for b in closing_data["bottlenecks"]:
            sev_badge = "🔴 CRÍTICO" if b["severity"] == "critical" else "🟡 SEGUIMIENTO"
            b_vals = [
                sev_badge,
                b["client_name"],
                b["resp"],
                b["stage"],
                f"{b['days_waiting']} días",
                f"${b['amount']:,.0f}".replace(",", ".") if b['amount'] else "$ 0",
                b["details"]
            ]
            for c_idx, v in enumerate(b_vals, start=1):
                cell = ws1.cell(row=curr_row, column=c_idx, value=v)
                cell.font = Font(name="Calibri", size=9, bold=(c_idx in [1, 5]))
                cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 4, 5] else "left", vertical="center")
                cell.border = border_thin
                if b["severity"] == "critical":
                    cell.fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
            curr_row += 1
    else:
        ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=7)
        c = ws1.cell(row=curr_row, column=1, value="✓ ¡Excelente! No hay cuellos de botella ni alertas críticas en este mes.")
        c.font = Font(name="Calibri", size=10, bold=True, color="059669")
        c.alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # --- SHEET 2: CONCILIACIÓN DETALLADA ---
    ws2 = wb.create_sheet(title="Conciliación Detallada")
    ws2.merge_cells("A1:Q1")
    ws2["A1"] = f"DETALLE DE FACTURACIÓN, PREFACTURAS Y ÓRDENES - {month.upper()}"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws2["A1"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30

    det_headers = [
        "ID", "CLIENTE", "RESPONSABLE", "FRECUENCIA", "PERÍODO",
        "P1: INFORME", "P1: FECHA", "Nº PREFACTURA", "VALOR PREFACTURA",
        "P2: OC/HES", "Nº OC", "P3: FACTURA", "FECHA FACTURA", "Nº FACTURA", "VALOR FACTURA",
        "P4: RADICADO", "ESTADO CIERRE"
    ]
    for c_idx, h_text in enumerate(det_headers, start=1):
        c = ws2.cell(row=3, column=c_idx, value=h_text)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    conn = get_db_connection()
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    if month and month != "ALL":
        cursor.execute("SELECT * FROM billing_records WHERE month = ? ORDER BY client_id ASC", (month,))
    else:
        cursor.execute("SELECT * FROM billing_records ORDER BY client_id ASC")
    rec_list = cursor.fetchall()
    cursor.execute("SELECT * FROM clients")
    cli_map = {c["id"]: dict(c) for c in cursor.fetchall()}
    conn.close()

    for r_idx, r in enumerate(rec_list, start=4):
        c_meta = cli_map.get(r["client_id"], {})
        status_label = "🟢 COMPLETADO" if r["step4"] else ("🔵 FACTURADO" if r["step3"] else ("🟡 ESP. OC" if r["step1"] else "⚪ PENDIENTE"))
        r_vals = [
            r["client_id"],
            r["client_name"],
            c_meta.get("resp", ""),
            r["freq_type"],
            r["period_detail"],
            "SÍ" if r["step1"] else "NO",
            r["step1_date"],
            r["pref_num"],
            r["pref_val"],
            "SÍ" if r["step2"] else "NO",
            r["step2_oc"],
            "SÍ" if r["step3"] else "NO",
            r["step3_date"],
            r["step3_fac"],
            r["fac_val"],
            "SÍ" if r["step4"] else "NO",
            status_label
        ]
        for c_idx, v in enumerate(r_vals, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=v)
            cell.font = Font(name="Calibri", size=9)
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2] else "left", vertical="center")
            cell.border = border_thin
            if r["step4"]:
                cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Informe_Cierre_Oficial_{month}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Serve Frontend SPA
@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"))

# Serve Uploaded Documents (with Disaster Recovery from DB Archive)
@app.get("/uploads/{filename}")
def serve_uploaded_file(filename: str):
    file_path = os.path.join(UPLOADS_DIR, filename)
    if os.path.exists(file_path):
        return FileResponse(file_path)

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT file_bytes, content_type, filename FROM uploaded_files_archive WHERE filename = ? OR file_url = ? OR file_url LIKE ?", (filename, f"/uploads/{filename}", f"%{filename}%"))
        row = cursor.fetchone()
        conn.close()

        if row and row[0]:
            raw_bytes = row[0]
            file_bytes = bytes(raw_bytes) if isinstance(raw_bytes, memoryview) else raw_bytes
            content_type = row[1] or "application/octet-stream"
            try:
                with open(file_path, "wb") as f:
                    f.write(file_bytes)
            except Exception:
                pass
            return Response(
                content=file_bytes,
                media_type=content_type,
                headers={"Content-Disposition": f"inline; filename={filename}"}
            )
    except Exception as e:
        print(f"Error recuperando archivo de la base de datos: {e}")

    raise HTTPException(status_code=404, detail="Archivo no encontrado")

app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")
app.mount("/", StaticFiles(directory=STATIC_DIR), name="static")

if __name__ == "__main__":
    import uvicorn
    print("Iniciando servidor de facturación en http://0.0.0.0:8000 ...")
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=False)
